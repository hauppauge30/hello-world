import tkinter  
from tkinter import *
window = tkinter.Tk()
import xlsxwriter
from docx import Document
import docx
import os


document = Document()
document.add_heading('Calibration Table DCDC_Ford_mHEV-C', 0)
document.save('D:\Calibration_Table_DCDC_Ford_mHEV-C.docx')

entry_text = Entry(window, width=60, bg='white',
                   textvariable='entry_text')

entry_text.grid(row = 2, column = 0,sticky = S)

entry_text_r = Entry(window, width=60, bg='white',
                     textvariable='entry_text_r')

entry_text_r.grid(row = 5, column = 0,sticky = S)



number_of_files_found = 0

window.geometry('486x250')
window.title("CCRG - Continental Calibration Report Generator")
window.configure(background = "orange");

window.resizable(width=False, height=False)

#create label
Label(window,text = "     Please enter the directory path for C_reports:",bg = "orange",fg = "black",font ="none 15 bold") . grid(row = 1, column = 0 , sticky = S)
Label(window,text = "Please enter the directory path to grl files:",bg = "orange",fg = "black",font ="none 15 bold") . grid(row = 4, column = 0 , sticky = S)
Label(window,text = "",bg = "orange",fg = "black",font ="none 15 bold") . grid(row = 6, column = 0 , sticky = S)
Button(window,text = "GENERATE REPORTS",width = 33 ,command = "click") .grid(row=7,column=0,sticky=S)
Label(window,text = "",bg = "orange",fg = "black",font ="none 15 bold") . grid(row = 8, column = 0 , sticky = S)

filename = ""

def click():
    #ASW
    number_of_files_found = Parse_File(entry_text_r.get() + "\work\ASW\EE\EDCMMG_PRJ\edcmmg_acq\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\ASW\EE\EDCMMG_PRJ\edcmmg_com\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\ASW\EE\EDCMMG_PRJ\edcmmg_fault_mng\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\ASW\EE\EDCMMG_PRJ\edcmmg_prj\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\ASW\EE\EDCPPR_PRJ\edcppr_cur\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\ASW\EE\EDCPPR_PRJ\edcppr_hw\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\ASW\EE\EDCPPR_PRJ\edcppr_prj\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\ASW\EE\EDCPPR_PRJ\edcppr_rteif\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\ASW\EE\EDCPPR_PRJ\edcppr_volt\src")
    #BSW
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\CDD\ECOMMG_PRJ\ecommg_mcan\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\CDD\ECOMMG_PRJ\ecommg_prj\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\CDD\ECOMMG_PRJ\ecommg_sbc\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\CDD\EDCCCT_PRJ\edccct_clamp\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\CDD\EDCCCT_PRJ\edccct_ctl\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\CDD\EDCCCT_PRJ\edccct_out\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\CDD\EDCCCT_PRJ\edccct_prj\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\CDD\EDCCCT_PRJ\edccct_rteif\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\CDD\EDCCCT_PRJ\edccct_sig\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\CDD\EPWMCT_PRJ\epwmct_drv\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\CDD\EPWMCT_PRJ\epwmct_prj\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\CDD\EPWMCT_PRJ\epwmct_drv\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\ECAL\IOHWABSTR_PRJ\iohwabstr_adc\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\ECAL\IOHWABSTR_PRJ\iohwabstr_dip\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\ECAL\IOHWABSTR_PRJ\iohwabstr_dop\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\ECAL\IOHWABSTR_PRJ\iohwabstr_prj\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\ECAL\IOHWABSTR_PRJ\iohwabstr_pwm\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\ECAL\IOHWABSTR_PRJ\iohwabstr_rng\src")
    number_of_files_found = number_of_files_found + Parse_File(entry_text_r.get() + "\work\BSW\ECAL\IOHWABSTR_PRJ\iohwabstr_rteif\src")
      
    if(number_of_files_found <= 0):
        NoFilesFound()
    else:
       Label(window,text = "The reports has been generated successfully",bg = "orange",fg = "green",font ="none 15 bold") . grid(row = 9, column = 0 , sticky = S) 
    
    print(number_of_files_found)




def NoFilesFound():
    Label(window,text = "                                                                              ",bg = "orange",fg = "red",font ="none 15 bold") . grid(row = 9, column = 0 , sticky = S)
    Label(window,text = " No files found ! ",bg = "orange",fg = "red",font ="none 15 bold") . grid(row = 9, column = 0 , sticky = S)

def Parse_File(directory_path): 
    i = 0    
    if(entry_text_r.get() == "" ):
        Label(window,text = "  Please enter paths ",bg = "orange",fg = "red",font ="none 15 bold") . grid(row = 9, column = 0 , sticky = S)
    os.chdir(directory_path)
    for file in os.listdir(directory_path):
        if file.endswith(".grl"):
            i = i +1   
            #print(os.path.join("", file))  
            filename = file;   
            Generate_Report(filename);
    
    return i 
              

def Generate_Report(name_of_file):
    #Excell Setting#
    Excell_name = name_of_file[:-4]
    workbook = xlsxwriter.Workbook(entry_text.get() + '\C_'+Excell_name+'_report.xlsx')
    worksheet = workbook.add_worksheet()

    text_format = workbook.add_format({'text_wrap': True,'bold': True, 'font_color': 'black', 'font_size': 16,'border':1})
    text_format.set_bg_color('lime')

    number_format = workbook.add_format({'text_wrap': True,'bold': True,'border':1,'font_size': 13})
    number_format.set_bg_color('white')
                               #number_format.set_num_format('#,##0') 

    #Excell Setting#                          
    width_Cal= len("Calibration Name") +35
    width_Val= len("Calibration Value") + 36
    width_Desc= len("Calibration Description") + 70
    width_Rem= len("Calibration Description") + 30
    worksheet.set_column(1, 1, width_Cal)
    worksheet.set_column(0, 1, width_Val)
    worksheet.set_column(2, 1, width_Desc)
    worksheet.set_column(3, 1, width_Rem)
    #Write some data headers.
    worksheet.write('A1', 'Calibration Name', text_format)
    worksheet.write('B1', 'Calibration Value', text_format)
    worksheet.write('C1', 'Calibration Description', text_format)
    worksheet.write('D1', 'Remarks', text_format)
    row_nam = 1
    row_val = 1
    row_desc = 1
    #Excell Setting#

    store_string = ""
    increment_C = 0
    vector_flag = 0

    FileName = (name_of_file)
    #This is set to 1 in case a first parameter was stored
    parameter_reached = 0
    description = []
    j = 1
    with open(FileName, 'r') as file:
        lines_in_file = file.read()
        #print  ("Number of Words: ", len(lines_in_file.split()))
        for i in range (0,len(lines_in_file.split())):
            if(lines_in_file.split()[i] == "parameter" and (lines_in_file.split()[i-1] == "{" or lines_in_file.split()[i-1] == "}") and (lines_in_file.split()[i-3] != "cTag") and (lines_in_file.split()[i-15] != "cTag") ):
                parameter_reached = 1
                #Print the name of Calibration
                #print(lines_in_file.split()[i+9])
                #print(lines_in_file.split()[i+1])
                if(lines_in_file.split()[i+1][:1] != "c" and lines_in_file.split()[i+1][:1] != "f" and lines_in_file.split()[i+1][:1] != "C" ):
                    #print(lines_in_file.split()[i+1][:1])
                    increment_C = increment_C +1
                    #print(increment_C)
            
                    
                worksheet.write(row_nam,0, lines_in_file.split()[i+1],number_format)
                row_nam = row_nam +1 
            
                #Print the value of Calibration
                #print(lines_in_file.split()[i+9])
                st = lines_in_file.split()[i+9][:-1]
                worksheet.write(row_val,1, st,number_format)
                worksheet.write(row_val,3, " ",number_format)
                row_val = row_val +1
            #else:
                #parameter_reached = 0  
                
                   
            if(parameter_reached == 1):    
                if(lines_in_file.split()[i] == "mlString" and (lines_in_file.split()[i-8] == "}" or lines_in_file.split()[i-8] == "mappingSchemeParameter" or lines_in_file.split()[i-4] == "mappingSchemeParameter" )):
                    #print(lines_in_file.split()[i-8])
                    i = i + 4;
                    while(lines_in_file.split()[i] != "language"):
                        if(vector_flag == 1):
                            while(increment_C != 0):
                                worksheet.write(row_desc,2, " ",number_format)
                                row_desc = row_desc + 1 
                                increment_C = increment_C - 1
                        
                        if(lines_in_file.split()[i] == "}"):
                            vector_flag = 1 
                            break
                        else:
                            vector_flag = 0
                        
                        
                        #Print the description of Calibration
                        description.append(lines_in_file.split()[i]) 
                        #print(lines_in_file.split()[i])
                        j = j +1
                        i = i+1
                    for i in range (len(description)):    
                        #print(description[i],end =" ")
                        store_string = store_string + " " + description[i]
                   
                    
                    description = []
                    #print(store_string)
                    store_string = store_string[:-1]
                    worksheet.write(row_desc,2, store_string,number_format)
                    row_desc = row_desc + 1 
                    store_string = ""  
            
        workbook.close()
        
window.mainloop()
